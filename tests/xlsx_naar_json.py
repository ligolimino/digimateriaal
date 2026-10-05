"""Zet de Excel-overzichtslijst om naar JSON voor de testomgeving (nabootsing van Google Sheets)."""
import json, sys, datetime, warnings
import openpyxl
warnings.filterwarnings('ignore')

bron, doel = sys.argv[1], sys.argv[2]
waarden_wb = openpyxl.load_workbook(bron, data_only=True)
formules_wb = openpyxl.load_workbook(bron)

def toon(v):
    """Ongeveer zoals Google Sheets een waarde toont (getDisplayValues)."""
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, datetime.datetime):
        return v.strftime('%d/%m/%Y')
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)

uit = []
for ws_w in waarden_wb.worksheets:
    ws_f = formules_wb[ws_w.title]
    laatste_rij, laatste_kol = 0, 0
    for rij in ws_f.iter_rows():
        for c in rij:
            if c.value not in (None, '') or c.hyperlink:
                laatste_rij = max(laatste_rij, c.row)
                laatste_kol = max(laatste_kol, c.column)
    tonen, links, formules = [], [], []
    for r in range(1, laatste_rij + 1):
        t, l, f = [], [], []
        for k in range(1, laatste_kol + 1):
            cf = ws_f.cell(r, k)
            t.append(toon(ws_w.cell(r, k).value))
            l.append(cf.hyperlink.target if cf.hyperlink and cf.hyperlink.target else '')
            f.append(cf.value if isinstance(cf.value, str) and cf.value.startswith('=') else '')
        tonen.append(t); links.append(l); formules.append(f)
    verborgen = [k for k in range(1, laatste_kol + 1)
                 if ws_f.column_dimensions[openpyxl.utils.get_column_letter(k)].hidden]
    uit.append({'naam': ws_w.title, 'verborgen': ws_w.sheet_state != 'visible',
                'verborgenKolommen': verborgen, 'tonen': tonen, 'links': links, 'formules': formules})
json.dump(uit, open(doel, 'w', encoding='utf8'), ensure_ascii=False)
print('\n'.join(f"{b['naam']}: {len(b['tonen'])} rijen x {len(b['tonen'][0]) if b['tonen'] else 0} kolommen" for b in uit))
